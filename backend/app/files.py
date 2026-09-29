"""
Fichiers joints : dépôt, extraction du texte des documents, profil des
fichiers de données, suppression.
"""
import asyncio
import io
import logging
import re
import shutil
import unicodedata
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import config, sandbox
from .auth import CurrentUser, current_user
from .db import get_db
from .models import File, new_id, now_ms

log = logging.getLogger("opti.files")
router = APIRouter(prefix="/api/files", tags=["files"])

DOCUMENT_EXT = {".pdf", ".docx", ".txt", ".md"}
DATA_EXT = {".csv", ".xlsx", ".xlsm"}
ALLOWED = DOCUMENT_EXT | DATA_EXT


def file_dir(file_id: str) -> Path:
    return config.FILES_DIR / file_id


def file_path(f: File) -> Path:
    return file_dir(f.id) / f.stored_name


def safe_name(filename: str) -> str:
    name = unicodedata.normalize("NFKD", filename).encode("ascii", "ignore").decode()
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._") or "fichier"
    stem, dot, ext = name.rpartition(".")
    return (stem[:60] + dot + ext.lower()) if dot else name[:60]


def file_summary(f: File) -> dict:
    return {"id": f.id, "filename": f.filename, "kind": f.kind, "size": f.size}


# ── Extraction du texte ──────────────────────────────────────────────────────
def extract_text(path: Path, ext: str) -> str:
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            txt = (page.extract_text() or "").strip()
            if txt:
                pages.append(f"[Page {i}]\n{txt}")
        return "\n\n".join(pages)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return ""


# ── Routes ───────────────────────────────────────────────────────────────────
@router.post("", status_code=201)
async def upload(file: UploadFile, user: CurrentUser = Depends(current_user), db: Session = Depends(get_db)):
    filename = (file.filename or "fichier").strip()[:255]
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(415, "Format non pris en charge. Formats acceptés : PDF, Word (.docx), texte, CSV, Excel (.xlsx).")

    content = await file.read(config.MAX_UPLOAD_MB * 1024 * 1024 + 1)
    if len(content) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(413, f"Fichier trop volumineux (maximum {config.MAX_UPLOAD_MB} Mo).")
    if not content:
        raise HTTPException(400, "Le fichier est vide.")

    f = File(id=new_id(), owner_id=user.id, filename=filename, stored_name=safe_name(filename),
             kind="data" if ext in DATA_EXT else "document", size=len(content))
    d = file_dir(f.id)
    d.mkdir(parents=True, exist_ok=True)
    d.chmod(0o755)
    path = d / f.stored_name
    path.write_bytes(content)
    path.chmod(0o644)   # lisible par l'utilisateur non privilégié du bac à sable

    try:
        if f.kind == "document":
            text = await asyncio.to_thread(extract_text, path, ext)
            if not text.strip():
                raise HTTPException(422, "Aucun texte n'a pu être extrait (document scanné ou protégé ?).")
            f.text = text
        else:
            f.text = await sandbox.profile_data_file(path, f.stored_name)
    except HTTPException:
        shutil.rmtree(d, ignore_errors=True)
        raise
    except Exception:
        log.exception("Lecture impossible de %s", filename)
        shutil.rmtree(d, ignore_errors=True)
        raise HTTPException(422, "Le fichier n'a pas pu être lu.")

    db.add(f)
    db.commit()
    return file_summary(f)


@router.delete("/{file_id}", status_code=204)
def remove(file_id: str, user: CurrentUser = Depends(current_user), db: Session = Depends(get_db)):
    f = db.get(File, file_id)
    if not f or f.owner_id != user.id:
        raise HTTPException(404, "Fichier introuvable")
    if f.message_id:
        raise HTTPException(409, "Ce fichier fait déjà partie d'une conversation.")
    db.delete(f)
    db.commit()
    shutil.rmtree(file_dir(file_id), ignore_errors=True)


def delete_files_of_chat(db: Session, chat_id: str) -> None:
    for f in db.scalars(select(File).where(File.chat_id == chat_id)):
        shutil.rmtree(file_dir(f.id), ignore_errors=True)


def purge_orphans(db: Session) -> int:
    """Supprime les fichiers déposés mais jamais envoyés depuis plus de 24 h."""
    limit = now_ms() - 86_400_000
    orphans = db.scalars(select(File).where(File.message_id.is_(None), File.created_at < limit)).all()
    for f in orphans:
        shutil.rmtree(file_dir(f.id), ignore_errors=True)
        db.delete(f)
    db.commit()
    return len(orphans)
